"""
Motor de Geração e Preenchimento de Planilhas SGT3 (Claro / Grupo Talentos).
Manipulação via openpyxl com suporte obrigatório a .xlsx e .xlsm (preservando macros com keep_vba=True).
Gera planilhas a partir de templates existentes ou do zero com os layouts oficiais rigorosamente respeitados.
"""

import os
import io
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


TEMPLATES_DIR = os.path.join(os.path.dirname(__file__), "templates")


def _get_or_create_workbook(template_bytes: Optional[bytes], default_filename: str, is_macro: bool = False):
    """
    Carrega o template do usuário em memória, ou da pasta templates/,
    ou cria um novo Workbook do zero se nenhum estiver disponível.
    """
    if template_bytes:
        return openpyxl.load_workbook(io.BytesIO(template_bytes), keep_vba=is_macro)
        
    local_path = os.path.join(TEMPLATES_DIR, default_filename)
    if os.path.exists(local_path):
        return openpyxl.load_workbook(local_path, keep_vba=is_macro)
        
    # Cria novo workbook
    wb = openpyxl.Workbook()
    return wb


def generate_criacao_identidade_excel(records: List[Dict[str, Any]], template_bytes: Optional[bytes] = None) -> bytes:
    """
    Preenche a planilha SGT3_CRIAÇÃO_DE_IDENTIDADE.xlsm.
    - Preserva macros se template for fornecido.
    - Dados iniciam na Linha 6.
    - Preenche colunas operacionais e valores institucionais fixos do Grupo Talentos / Claro.
    """
    is_macro = True
    wb = _get_or_create_workbook(template_bytes, "SGT3_CRIAÇÃO_DE_IDENTIDADE.xlsm", is_macro=is_macro)
    ws = wb.active
    ws.title = "Identidades SGT3"

    # Se criamos do zero, adiciona cabeçalhos operacionais até a linha 5
    if not template_bytes and not os.path.exists(os.path.join(TEMPLATES_DIR, "SGT3_CRIAÇÃO_DE_IDENTIDADE.xlsm")):
        _build_criacao_identidade_header(ws)

    # Inserção de dados a partir da Linha 6
    start_row = 6
    for idx, rec in enumerate(records):
        row_num = start_row + idx
        
        # Colunas fixas e variáveis conforme especificação SGT3
        ws.cell(row=row_num, column=2, value="TERCEIRO CALL CENTER")               # B (2)
        ws.cell(row=row_num, column=3, value="NÃO")                                # C (3)
        ws.cell(row=row_num, column=4, value="SIM (CLARO + NET + EMBRATEL)")       # D (4)
        ws.cell(row=row_num, column=5, value="NÃO")                                # E (5)
        ws.cell(row=row_num, column=6, value="NÃO SE APLICA")                      # F (6)
        
        # Dados do colaborador
        ws.cell(row=row_num, column=9, value=rec.get("cpf_formatado", ""))          # I (9)
        ws.cell(row=row_num, column=10, value=rec.get("nome_completo", ""))        # J (10)
        ws.cell(row=row_num, column=11, value=rec.get("data_nascimento", ""))      # K (11)
        ws.cell(row=row_num, column=13, value=rec.get("telefone_celular", ""))     # M (13)
        ws.cell(row=row_num, column=14, value=rec.get("email", "suporte@grupotalentos.com.br")) # N (14)
        ws.cell(row=row_num, column=15, value=rec.get("qualificacao", "MÉDIO"))    # O (15)
        
        # Dados institucionais Talentos
        ws.cell(row=row_num, column=16, value="54.956.511/0001-00")                # P (16)
        ws.cell(row=row_num, column=17, value="Talentos Total Servicos LTDA ME")   # Q (17)
        ws.cell(row=row_num, column=18, value="CALL CENTER - COBRANÇA")            # R (18)
        ws.cell(row=row_num, column=19, value="NÃO")                               # S (19)
        ws.cell(row=row_num, column=20, value="54.956.511/0001-00")                # T (20)
        ws.cell(row=row_num, column=21, value="Talentos Total Servicos LTDA ME")   # U (21)
        ws.cell(row=row_num, column=22, value="CLT")                               # V (22)
        
        # Hierarquia e Gestão
        ws.cell(row=row_num, column=25, value="SUPORTE AO NEGOCIO")                # Y (25)
        ws.cell(row=row_num, column=26, value="TERCEIRIZAÇÃO")                     # Z (26)
        ws.cell(row=row_num, column=27, value="SERVIÇOS DE APOIO")                 # AA (27)
        ws.cell(row=row_num, column=28, value="CLARO SP - REGIONAL SÃO PAULO")     # AB (28)
        ws.cell(row=row_num, column=29, value="SP")                                # AC (29)
        ws.cell(row=row_num, column=30, value="São Paulo")                         # AD (30)
        ws.cell(row=row_num, column=31, value="028.463.376-32")                    # AE (31)
        ws.cell(row=row_num, column=32, value="DEGLEITON AQUINO DE MELO")          # AF (32)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _build_criacao_identidade_header(ws):
    """Constrói cabeçalho padrão elegante para Criação de Identidade."""
    ws.merge_cells("A1:AF3")
    title_cell = ws["A1"]
    title_cell.value = "SGT3 - CRIAÇÃO DE IDENTIDADE / SOLICITAÇÃO DE ACESSOS (V 16)"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="DA291C", end_color="DA291C", fill_type="solid") # Claro Red
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    headers_row5 = {
        2: "TIPO TERCEIRO", 3: "VIP", 4: "EMPRESAS", 5: "ESTRANGEIRO", 6: "TIPO DOC",
        9: "CPF", 10: "NOME COMPLETO", 11: "DT NASCIMENTO", 13: "CELULAR",
        14: "E-MAIL", 15: "QUALIFICAÇÃO", 16: "CNPJ CONTRATANTE", 17: "RAZÃO SOCIAL CONTRATANTE",
        18: "ÁREA DE ATUAÇÃO", 19: "HOMOLOGAÇÃO", 20: "CNPJ FORNECEDOR", 21: "RAZÃO SOCIAL FORNECEDOR",
        22: "TIPO CONTRATO", 25: "DIRETORIA", 26: "GERÊNCIA", 27: "COORDENAÇÃO",
        28: "ESTRUTURA LOCAL", 29: "UF", 30: "MUNICÍPIO", 31: "CPF GESTOR", 32: "NOME GESTOR"
    }
    for col_idx, text in headers_row5.items():
        c = ws.cell(row=5, column=col_idx, value=text)
        c.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def generate_revogacao_excel(records: List[Dict[str, Any]], template_bytes: Optional[bytes] = None) -> bytes:
    """
    Preenche a planilha SGT3_REVOGAÇÂO.xlsx.
    - Linha 1 = Cabeçalho: LOGIN Z, NOME, CPF, EMPRESA
    - Linha 2 em diante = Dados
    """
    wb = _get_or_create_workbook(template_bytes, "SGT3_REVOGAÇÂO.xlsx", is_macro=False)
    ws = wb.active
    ws.title = "Revogação"

    # Cabeçalhos na linha 1
    headers = ["LOGIN Z", "NOME", "CPF", "EMPRESA"]
    header_fill = PatternFill(start_color="DA291C", end_color="DA291C", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, rec in enumerate(records, start=2):
        ws.cell(row=idx, column=1, value=rec.get("login_z", "")).alignment = Alignment(horizontal="center")
        ws.cell(row=idx, column=2, value=rec.get("nome", ""))
        # CPF gravado estritamente como texto com 11 dígitos
        c_cell = ws.cell(row=idx, column=3, value=str(rec.get("cpf_numeros", "")))
        c_cell.number_format = "@"
        c_cell.alignment = Alignment(horizontal="center")
        ws.cell(row=idx, column=4, value=rec.get("empresa", "Talentos Total Servicos LTDA ME"))

    # Auto ajuste de colunas
    for col in ws.columns:
        max_len = max(len(str(c.value or '')) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def generate_atualiza_status_excel(records: List[Dict[str, Any]], template_bytes: Optional[bytes] = None) -> bytes:
    """
    Preenche a planilha SGT3_ATUALIZA_STATUS.xlsm.
    - Linha 1 = Ação / Título
    - Linha 2 = Cabeçalhos
    - Linha 3 em diante = Dados
    """
    wb = _get_or_create_workbook(template_bytes, "SGT3_ATUALIZA_STATUS.xlsm", is_macro=True)
    ws = wb.active
    ws.title = "Status Colaboradores"

    # Linha 1 - Ação
    ws.merge_cells("A1:H1")
    t1 = ws["A1"]
    t1.value = "SGT3 - ATUALIZAÇÃO DE STATUS OPERACIONAL"
    t1.font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    t1.fill = PatternFill(start_color="DA291C", end_color="DA291C", fill_type="solid")
    t1.alignment = Alignment(horizontal="center", vertical="center")

    # Linha 2 - Cabeçalhos
    headers = [
        (1, "CPF"),
        (2, ""), # Em branco
        (3, "NOME COMPLETO"),
        (4, "RAZÃO SOCIAL"),
        (5, "FERIAS"),
        (6, "RETORNO FERIAS"),
        (7, "LICENCA"),
        (8, "RETORNO LICENCA"),
    ]
    h_fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
    h_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")

    for col_idx, h_text in headers:
        c = ws.cell(row=2, column=col_idx, value=h_text)
        c.fill = h_fill
        c.font = h_font
        c.alignment = Alignment(horizontal="center", vertical="center")

    # Linha 3 em diante - Dados
    for idx, rec in enumerate(records, start=3):
        # Coluna A: CPF 11 dígitos texto
        cpf_cell = ws.cell(row=idx, column=1, value=str(rec.get("cpf_numeros", "")))
        cpf_cell.number_format = "@"
        cpf_cell.alignment = Alignment(horizontal="center")
        
        # Coluna B: Em branco
        ws.cell(row=idx, column=2, value="")
        
        # Coluna C: Nome Completo
        ws.cell(row=idx, column=3, value=rec.get("nome_completo", ""))
        
        # Coluna D: Razão Social fixa
        ws.cell(row=idx, column=4, value="Talentos Total Servicos LTDA ME")
        
        # Colunas E, F, G, H: X para a respectiva situação
        sit = rec.get("situacao", "")
        ws.cell(row=idx, column=5, value="X" if sit == "FERIAS" else "").alignment = Alignment(horizontal="center")
        ws.cell(row=idx, column=6, value="X" if sit == "RETORNO FERIAS" else "").alignment = Alignment(horizontal="center")
        ws.cell(row=idx, column=7, value="X" if sit == "LICENCA" else "").alignment = Alignment(horizontal="center")
        ws.cell(row=idx, column=8, value="X" if sit == "RETORNO LICENCA" else "").alignment = Alignment(horizontal="center")

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def generate_duo_desktop_excel(records: List[Dict[str, Any]], template_bytes: Optional[bytes] = None) -> bytes:
    """
    Preenche a planilha SGT3_DUO_DESKTOP.xlsx.
    - Dados iniciam estritamente na Linha 8.
    - Coluna D (4): CLARO BR
    - Coluna L (12): Login Z
    - Coluna N (14): Nome Completo
    - Coluna R (18): Hostname
    - Coluna T (20): Machineguid (em minúsculas)
    - Demais colunas em branco.
    """
    wb = _get_or_create_workbook(template_bytes, "SGT3_DUO_DESKTOP.xlsx", is_macro=False)
    ws = wb.active
    ws.title = "DUO Desktop"

    if not template_bytes and not os.path.exists(os.path.join(TEMPLATES_DIR, "SGT3_DUO_DESKTOP.xlsx")):
        # Preenche cabeçalho estético até linha 7
        ws.merge_cells("A1:U4")
        t = ws["A1"]
        t.value = "SGT3 - ATIVAÇÃO DUO DESKTOP / TOKEN CISCO & CLARO"
        t.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        t.fill = PatternFill(start_color="DA291C", end_color="DA291C", fill_type="solid")
        t.alignment = Alignment(horizontal="center", vertical="center")

        ws.cell(row=7, column=4, value="ORIGEM").font = Font(bold=True)
        ws.cell(row=7, column=12, value="LOGIN Z").font = Font(bold=True)
        ws.cell(row=7, column=14, value="NOME COMPLETO").font = Font(bold=True)
        ws.cell(row=7, column=18, value="HOSTNAME").font = Font(bold=True)
        ws.cell(row=7, column=20, value="MACHINEGUID").font = Font(bold=True)

    # Inicia estritamente na Linha 8
    start_row = 8
    for idx, rec in enumerate(records):
        row_num = start_row + idx
        ws.cell(row=row_num, column=4, value="CLARO BR").alignment = Alignment(horizontal="center")
        ws.cell(row=row_num, column=12, value=rec.get("login_z", "")).alignment = Alignment(horizontal="center")
        ws.cell(row=row_num, column=14, value=rec.get("nome_completo", ""))
        ws.cell(row=row_num, column=18, value=rec.get("hostname", "")).alignment = Alignment(horizontal="center")
        ws.cell(row=row_num, column=20, value=str(rec.get("machineguid", "")).lower()).alignment = Alignment(horizontal="center")

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
